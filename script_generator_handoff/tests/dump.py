import openpyxl, sys
for f in sys.argv[1:]:
    wb = openpyxl.load_workbook(f, data_only=False)
    print("=====", f, wb.sheetnames)
    for ws in wb.worksheets:
        print(f"--- sheet '{ws.title}' dims={ws.dimensions} merged={list(ws.merged_cells.ranges)[:5]} dv={[ (d.sqref, d.formula1) for d in ws.data_validations.dataValidation][:5]}")
        for row in ws.iter_rows():
            vals = [(c.coordinate, c.value) for c in row if c.value is not None]
            if vals: print("  ", " | ".join(f"{k}={v!r}" for k,v in vals))
